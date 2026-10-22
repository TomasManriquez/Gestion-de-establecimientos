"""Migra `establishments.connectivity` al modelo en tres bloques: 2030, BAM y otro tipo.

Por defecto solo muestra la vista previa (no modifica nada):

    python scripts/migrate_connectivity.py --excel ../Documentacion/DASHBOARD_final.xlsx
    python scripts/migrate_connectivity.py --excel ... --report vista_previa.csv   # además, en CSV
    python scripts/migrate_connectivity.py --excel ... --apply                      # escribe en Mongo

Sin --excel convierte solo con lo que hay en la base (el bloque «otro tipo» queda sin tipo,
vencimiento ni OC). Con --from-json <archivo> lee el JSON de la carga inicial en vez de Mongo
(sirve para revisar sin base) y, con --write-json <salida>, deja ese JSON ya convertido para que
una instalación nueva siembre el modelo nuevo.

Es idempotente: un documento ya migrado no se toca. Usa MONGODB_URL y DATABASE_NAME del entorno.
No imprime claves de wifi. Hacer un respaldo antes de --apply (`scripts/backup-mongo.sh`).
"""
import argparse
import asyncio
import csv
import json
import os
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.database.connectivity_migration import (has_other, migrate_connectivity, rows_from_sheet,  # noqa: E402
                                                 without_internet)


def load_sheet(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb["Internet"].iter_rows(values_only=True)]
    return rows_from_sheet(rows)


def preview_rows(docs, sheet):
    out = []
    for doc in docs:
        rbd = str(doc.get("rbd", ""))
        new, notes = migrate_connectivity(doc.get("connectivity"), sheet.get(rbd) if sheet is not None else None)
        p, o = new["program_2030"], new["other"]
        out.append({
            "rbd": rbd, "nombre": doc.get("name", ""),
            "2030": f"{p['status'] or 'sin estado'} · {p['download_speed'] or '?'} Mbps" if p["enabled"] else "",
            "bam": len(new["bam"]) or "",
            "otro": " · ".join(x for x in (o["provider"], o["type"] or "", o["expiry_date"] and f"vence {o['expiry_date']}") if x)
                    if has_other(o) else "",
            "sin_internet": "sí" if without_internet(new) else "",
            "notas": "; ".join(notes),
            "_new": new, "_done": "ya migrado: sin cambios" in notes,
        })
    return out


def print_preview(rows, sheet, docs):
    w = {"rbd": 6, "nombre": 34, "2030": 26, "bam": 3, "otro": 40, "sin_internet": 4}
    head = f"{'RBD':<6} {'Establecimiento':<34} {'2030':<26} {'BAM':<3} {'Otro tipo':<40} {'Sin':<4} Notas"
    print(head)
    print("-" * len(head))
    for r in rows:
        print(" ".join(f"{str(r[k])[:n]:<{n}}" for k, n in w.items()), r["notas"])
    print()
    c = Counter()
    for r in rows:
        c["2030"] += bool(r["2030"]); c["BAM"] += bool(r["bam"]); c["otro tipo"] += bool(r["otro"])
        c["sin internet"] += bool(r["sin_internet"]); c["con notas para revisar"] += bool(r["notas"] and not r["_done"])
        c["ya migrados"] += r["_done"]
    print(f"Establecimientos: {len(rows)} · " + " · ".join(f"{k}: {v}" for k, v in c.items()))
    if sheet is not None:
        rbds = {str(d.get('rbd')) for d in docs}
        missing = sorted(set(sheet) - rbds)
        if missing:
            print(f"Filas de la planilla sin establecimiento en la base (se ignoran): {', '.join(missing)}")
        no_row = sorted(rbds - set(sheet))
        if no_row:
            print(f"Establecimientos sin fila en la planilla (se convierten solo con la base): {', '.join(no_row)}")


def write_report(rows, path):
    keys = ["rbd", "nombre", "2030", "bam", "otro", "sin_internet", "notas"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore", delimiter=";")
        wr.writeheader()
        wr.writerows(rows)
    print(f"Vista previa guardada en {path}")


async def main(args) -> int:
    sheet = load_sheet(args.excel) if args.excel else None
    if args.from_json:
        docs = json.loads(Path(args.from_json).read_text(encoding="utf-8"))
        col = None
    else:
        from motor.motor_asyncio import AsyncIOMotorClient
        client = AsyncIOMotorClient(os.environ.get("MONGODB_URL", "mongodb://localhost:27017"))
        col = client[os.environ.get("DATABASE_NAME", "slep_llanquihue")].establishments
        docs = [d async for d in col.find({}, {"rbd": 1, "name": 1, "connectivity": 1})]

    rows = preview_rows(docs, sheet)
    print_preview(rows, sheet, docs)
    if args.report:
        write_report(rows, args.report)

    if args.write_json:
        if not args.from_json:
            print("--write-json requiere --from-json")
            return 2
        by_rbd = {r["rbd"]: r["_new"] for r in rows}
        for d in docs:
            d["connectivity"] = by_rbd[str(d.get("rbd"))]
        Path(args.write_json).write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"JSON convertido escrito en {args.write_json}")

    if not args.apply:
        print("Modo vista previa: no se modificó nada. Use --apply para escribir en la base.")
        return 0
    if col is None:
        print("--apply escribe en Mongo; no se combina con --from-json")
        return 2
    changed = 0
    for r in rows:
        if r["_done"]:
            continue
        res = await col.update_one({"rbd": r["rbd"]}, {"$set": {"connectivity": r["_new"]}})
        changed += res.modified_count
    print(f"Documentos modificados: {changed}")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", help="DASHBOARD_final.xlsx (hoja «Internet»)")
    ap.add_argument("--report", help="guardar la vista previa en CSV")
    ap.add_argument("--from-json", help="leer el JSON de la carga inicial en vez de Mongo")
    ap.add_argument("--write-json", help="con --from-json: escribir el JSON convertido")
    ap.add_argument("--apply", action="store_true", help="escribir los cambios en Mongo")
    sys.exit(asyncio.run(main(ap.parse_args())))
